"""Celery tasks for Excel template filling pipeline.

Task Flow:
1. detect_fields_task: Analyze PDF to identify structured fields from Azure DI
2. auto_map_fields_task: Match PDF fields to Excel cells (LLM)
3. fill_excel_task: Fill Excel template with data (openpyxl)
4. start_fill_run_chain: Orchestrate the full pipeline
5. continue_fill_run_chain: Resume after user review
"""

import asyncio
import json
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

from celery import chain, shared_task
from celery.exceptions import SoftTimeLimitExceeded
from sqlalchemy.orm import Session

from app.config import settings
from app.database import SessionLocal
from app.repositories.document_repository import DocumentRepository
from app.repositories.job_repository import JobRepository
from app.repositories.template_repository import TemplateRepository
from app.services.beta_limits import commit_shadow_credits, reverse_shadow_credits
from app.services.citations import normalize_citations
from app.services.job_tracker import JobProgressTracker
from app.core.storage.storage_factory import get_storage_backend
from app.utils.logging import logger
from app.utils.metrics_recorder import record_template_fill_completed, record_template_fill_failed
from app.utils.metrics import TEMPLATE_FILL_LATENCY_SECONDS
from app.verticals.real_estate.template_filling.excel_handler import ExcelHandler
from app.verticals.real_estate.template_filling.llm_service import TemplateFillLLMService
from app.verticals.real_estate.template_filling.excel.mapping_coordinator import coordinator as mapping_coordinator
from app.verticals.real_estate.template_filling.source_map import (
    STRUCTURE_HIGH_CONFIDENCE,
    normalize_om_structure_with_pdf_fields,
)

# Domain modules
from app.verticals.real_estate.template_filling.artifacts import build_om_structure_artifact
from app.verticals.real_estate.template_filling.citations import (
    build_citation_context,
    get_field_page,
    get_section_citation_pages,
    resolve_bbox_from_citations,
)
from app.verticals.real_estate.template_filling.context_budget import build_budgeted_pdf_fields
from app.verticals.real_estate.template_filling.mapping_helpers import (
    build_narrative_pdf_field,
    build_scalar_context_for_batch,
    consolidate_scalar_batches_by_context,
    extract_schema_table_row_labels,
    get_scalar_batch_key,
    get_table_batch_key,
)
from app.verticals.real_estate.template_filling.schema_planner import (
    build_yaml_cell_status,
    compute_schema_counts,
    plan_schema_targets_for_structure,
)

# Backward-compat re-export used by tests
_plan_schema_targets_for_structure = plan_schema_targets_for_structure


# ---------------------------------------------------------------------------
# Infrastructure helpers
# ---------------------------------------------------------------------------

def _get_db_session() -> Session:
    return SessionLocal()


def _reverse_template_fill_shadow(fill_run_id: str, reason: str, stage: str) -> None:
    reverse_shadow_credits(
        operation_type="template_fill_run",
        reference_id=fill_run_id,
        reason=reason,
        metadata={"stage": stage},
    )


def _fail_fill_run(
    repo: TemplateRepository,
    tracker: JobProgressTracker,
    fill_run_id: str,
    stage: str,
    exc: Exception,
    *,
    is_retryable: bool = False,
) -> None:
    """Persist a stage failure to the DB and emit SSE error+end so the frontend terminates."""
    repo.update_fill_run(
        fill_run_id,
        status="failed",
        error_stage=stage,
        error_message=str(exc),
    )
    tracker.mark_error(
        error_stage=stage,
        error_message=str(exc),
        error_type=f"{stage}_failed",
        is_retryable=is_retryable,
    )


def _mark_auto_mapping_exception(
    repo: TemplateRepository,
    tracker: JobProgressTracker,
    fill_run_id: str,
    exc: Exception,
) -> None:
    _fail_fill_run(repo, tracker, fill_run_id, "auto_mapping", exc)


def _prepare_extracted_data_for_fill(
    extracted_data: Optional[Dict[str, Any]],
    field_mapping: Dict[str, Any],
) -> Dict[str, Any]:
    """Prepare extracted data for Excel filling without dropping stored artifacts."""

    prepared = dict(extracted_data) if isinstance(extracted_data, dict) else {}

    raw_llm_extracted = prepared.get("llm_extracted")
    if isinstance(raw_llm_extracted, dict):
        llm_extracted = dict(raw_llm_extracted)  # copy to avoid mutating caller's data
    else:
        if raw_llm_extracted is not None:
            logger.warning(
                "llm_extracted is not a dict (type: %s), resetting to empty dict",
                type(raw_llm_extracted),
            )
        llm_extracted = {}

    manual_edits = prepared.get("manual_edits", {})
    if not isinstance(manual_edits, dict):
        logger.warning(
            "manual_edits is not a dict (type: %s), resetting to empty dict",
            type(manual_edits),
        )
        manual_edits = {}
    else:
        cleaned_manual_edits = {}
        for sheet_name, cells in manual_edits.items():
            if isinstance(cells, dict):
                cleaned_manual_edits[sheet_name] = cells
            else:
                logger.warning(
                    "Skipping corrupted manual_edits entry: %s (type: %s)",
                    sheet_name,
                    type(cells),
                )
        manual_edits = cleaned_manual_edits

    for pdf_field in field_mapping.get("pdf_fields", []):
        field_id = pdf_field.get("id")
        auto_mapped_value = pdf_field.get("extracted_value")

        if field_id and field_id not in llm_extracted and auto_mapped_value is not None:
            field_entry = {
                "value": auto_mapped_value,
                "confidence": pdf_field.get("confidence", 0.95),
                "citations": pdf_field.get("citations", []),
                "user_edited": False,
            }
            if "bbox" in pdf_field:
                field_entry["bbox"] = pdf_field["bbox"]

            llm_extracted[field_id] = field_entry

    prepared["llm_extracted"] = llm_extracted
    prepared["manual_edits"] = manual_edits
    return prepared


def _build_targeted_virtual_pdf_field(
    mapping: Dict[str, Any],
    citation_context: Dict[str, Any],
) -> Dict[str, Any]:
    """Build a virtual pdf_field for a targeted schema mapping."""

    citations = normalize_citations(mapping.get("citations", []))
    virtual_field = {
        "id": mapping.get("pdf_field_id"),
        "name": mapping.get("pdf_field_name"),
        "type": mapping.get("data_type", "text"),
        "extracted_value": mapping.get("extracted_value"),
        "confidence": mapping.get("confidence", 0.7),
        "citations": citations,
        "reasoning": mapping.get("reasoning"),
        "source": "targeted_schema",
    }
    for display_key in ("display_label", "display_context", "source_note"):
        if mapping.get(display_key):
            virtual_field[display_key] = mapping[display_key]

    field_bbox = resolve_bbox_from_citations(citations, citation_context)
    if field_bbox:
        virtual_field["bbox"] = field_bbox

    return virtual_field


# ---------------------------------------------------------------------------
# Task: detect_fields_task
# ---------------------------------------------------------------------------

@shared_task(bind=True)
def detect_fields_task(self, payload: Dict[str, Any]) -> Dict[str, Any]:
    """
    Detect structured fields from PDF using Azure DI key-value pairs and tables.

    Uses free Azure DI extraction — no LLM cost.
    """
    fill_run_id = payload["fill_run_id"]
    document_id = payload["document_id"]
    job_id = payload["job_id"]

    db = _get_db_session()
    repo = TemplateRepository(db)
    tracker = JobProgressTracker(db, job_id)

    fill_run = repo.get_fill_run_with_data(fill_run_id)
    if not fill_run:
        raise ValueError(f"Fill run not found: {fill_run_id}")
    org_id = fill_run.org_id

    _DETECT_DONE_STATUSES = {"fields_detected", "mapping", "awaiting_review", "filling", "completed", "failed"}
    if fill_run.status in _DETECT_DONE_STATUSES:
        logger.warning(
            f"detect_fields_task skipped — fill_run={fill_run_id} already in status={fill_run.status!r}"
        )
        _idempotent_field_mapping = fill_run.fill_run_data.field_mapping if fill_run.fill_run_data else None
        db.close()
        return {**payload, "detection_result": {"fields": _idempotent_field_mapping.get("pdf_fields", []) if _idempotent_field_mapping else [], "total_fields": 0, "categories": []}}

    try:
        logger.info(f"Detecting fields for fill run: {fill_run_id}")

        tracker.update_progress(
            status="detecting_fields",
            current_stage="field_detection",
            progress_percent=20,
            message="Extracting fields from Azure DI key-value pairs and tables"
        )

        document_repo = DocumentRepository()
        document = document_repo.get_by_id(document_id, org_id)
        if not document:
            raise ValueError(f"Document not found: {document_id}")

        template = repo.get_template(fill_run.template_id)
        schema_cell_count = 0
        if template and template.schema_metadata and template.schema_metadata.get("schema_summary"):
            schema_cell_count = template.schema_metadata.get("total_yaml_fields")

        start_time = time.time()

        chunks = document_repo.get_chunks_for_document(document_id)
        if not chunks:
            raise ValueError(f"No chunks found for document {document_id}. Document may not be fully processed.")

        detected_fields = []
        field_id_counter = 1

        kv_chunks = []
        table_chunks = []
        narrative_chunks = []

        for chunk in chunks:
            metadata = chunk.chunk_metadata or {}
            if isinstance(metadata, str):
                try:
                    metadata = json.loads(metadata)
                except (json.JSONDecodeError, TypeError):
                    metadata = {}

            section_type = chunk.section_type or ""

            if section_type == "key_value_pairs":
                kv_chunks.append((chunk, metadata))
            elif section_type == "table":
                table_chunks.append((chunk, metadata))
            else:
                narrative_chunks.append((chunk, metadata))

        logger.info(f"Found {len(kv_chunks)} key-value chunks, {len(table_chunks)} table chunks, and {len(narrative_chunks)} narrative chunks")

        # Process KV chunks
        for chunk, metadata in kv_chunks:
            kv_pairs = metadata.get("key_value_pairs", [])

            for kv in kv_pairs:
                key = kv.get("key", "")
                value = kv.get("value", "")
                confidence = kv.get("confidence", 0.95)
                kv_bbox = kv.get("bbox", {})
                kv_page_number = (
                    (kv_bbox.get("page") if isinstance(kv_bbox, dict) and kv_bbox else None)
                    or kv.get("page_number")
                    or metadata.get("page_number")
                    or chunk.page_number
                )
                if isinstance(kv_page_number, str) and kv_page_number.isdigit():
                    kv_page_number = int(kv_page_number)
                citation = f"[S{field_id_counter}:p{kv_page_number}]"

                if not key:
                    continue

                field_data = {
                    "id": f"kv_{field_id_counter}",
                    "name": key,
                    "type": "text",
                    "extracted_value": value or "",
                    "confidence": confidence,
                    "citations": [citation],
                    "description": f"Key-value field from page {kv_page_number}",
                    "source": "key_value_pairs"
                }
                if kv_bbox:
                    field_data["bbox"] = kv_bbox

                detected_fields.append(field_data)
                field_id_counter += 1

        # Process table chunks
        for chunk, metadata in table_chunks:
            table_dict = None
            if chunk.tables:
                tables_data = chunk.tables
                if isinstance(tables_data, str):
                    try:
                        tables_data = json.loads(tables_data)
                    except (json.JSONDecodeError, TypeError):
                        tables_data = None

                if isinstance(tables_data, list) and len(tables_data) > 0:
                    try:
                        table_dict = tables_data[0]
                    except (IndexError, TypeError):
                        pass

            if not table_dict:
                table_dict = {
                    "table_name": metadata.get("table_name", ""),
                    "column_headers": metadata.get("column_headers", []),
                    "table_data": metadata.get("table_data", []),
                    "page_number": metadata.get("page_number"),
                }

            table_name = table_dict.get("table_name", "")
            column_headers = table_dict.get("column_headers", [])
            chunk_bbox = metadata.get("bbox", {})

            table_page = table_dict.get("page_number")
            if not table_page:
                page_range = metadata.get("page_range", [])
                metadata_page = metadata.get("page_number")
                range_page = page_range[0] if isinstance(page_range, list) and page_range else None
                table_page = (
                    (chunk_bbox.get("page") if isinstance(chunk_bbox, dict) and chunk_bbox else None)
                    or metadata_page
                    or range_page
                    or chunk.page_number
                )
            if isinstance(table_page, str) and table_page.isdigit():
                table_page = int(table_page)
            citation = f"[S{field_id_counter}:p{table_page}]"

            table_data_rows = table_dict.get("table_data", [])

            for col_idx, col_header in enumerate(column_headers):
                if not col_header or col_header.lower() in ["", "none", "n/a"]:
                    continue

                extracted_value = ""
                if table_data_rows and len(table_data_rows) > 0:
                    try:
                        if col_idx < len(table_data_rows[0]):
                            extracted_value = table_data_rows[0][col_idx]
                    except (ValueError, IndexError):
                        pass

                field_data = {
                    "id": f"tbl_{field_id_counter}",
                    "name": col_header,
                    "type": "text",
                    "extracted_value": extracted_value,
                    "confidence": 0.9,
                    "citations": [citation],
                    "description": f"Column from table '{table_name}' on page {table_page}",
                    "source": "table"
                }

                if chunk_bbox:
                    field_data["bbox"] = chunk_bbox

                detected_fields.append(field_data)
                field_id_counter += 1

            if column_headers:
                detected_fields.append({
                    "id": f"tbl_block_{field_id_counter}",
                    "name": table_name or "Table",
                    "type": "table",
                    "extracted_value": "",
                    "confidence": 0.9,
                    "citations": [citation],
                    "source": "table_block",
                    "table_name": table_name,
                    "table_columns": column_headers,
                    "table_rows": table_data_rows,
                    "page_number": table_page,
                })
                field_id_counter += 1

        # Process narrative chunks
        for chunk, metadata in narrative_chunks:
            narrative_text = chunk.text or ""
            if not narrative_text.strip():
                continue

            metadata_dict = metadata if isinstance(metadata, dict) else {}
            section_heading = ""
            if isinstance(metadata_dict.get("heading_hierarchy"), list):
                section_heading = metadata_dict["heading_hierarchy"][-1] if metadata_dict["heading_hierarchy"] else ""

            narrative_page = metadata_dict.get("page_number") or chunk.page_number
            if isinstance(narrative_page, str) and narrative_page.isdigit():
                narrative_page = int(narrative_page)

            detected_fields.append(
                build_narrative_pdf_field(
                    field_id_counter=field_id_counter,
                    narrative_text=narrative_text,
                    section_heading=section_heading,
                    narrative_page=narrative_page,
                )
            )
            field_id_counter += 1

        elapsed_ms = int((time.time() - start_time) * 1000)

        fields_with_bbox = sum(1 for f in detected_fields if "bbox" in f)
        kv_fields_with_bbox = sum(1 for f in detected_fields if f.get("source") == "key_value_pairs" and "bbox" in f)
        table_fields_with_bbox = sum(1 for f in detected_fields if f.get("source") == "table" and "bbox" in f)
        logger.info(
            f"[FIELD_DETECTION] {len(detected_fields)} total fields, "
            f"{fields_with_bbox} with bbox ({kv_fields_with_bbox} KV + {table_fields_with_bbox} table)"
        )

        detection_result = {
            "fields": detected_fields,
            "total_fields": len(detected_fields),
            "categories": ["key_value_pairs", "tables"]
        }

        field_mapping = {
            "pdf_fields": detected_fields,
            "mappings": []
        }

        citation_context = build_citation_context(detected_fields, document.filename)

        total_to_detect = schema_cell_count if schema_cell_count > 0 else len(detected_fields)

        repo.update_fill_run_data(fill_run_id, field_mapping=field_mapping, citation_context=citation_context)
        repo.update_fill_run(
            fill_run_id,
            total_fields_detected=total_to_detect,
            field_detection_completed=True,
            status="fields_detected",
            current_stage="field_detection",
        )

        db.close()

        logger.info(
            f"Field detection complete: {len(detected_fields)} fields detected "
            f"({len([f for f in detected_fields if f['source'] == 'key_value_pairs'])} from KV, "
            f"{len([f for f in detected_fields if f['source'] == 'table'])} from tables) "
            f"in {elapsed_ms}ms (Azure DI extraction - no LLM cost!)"
        )

        tracker.update_progress(
            status="fields_detected",
            current_stage="field_detection",
            progress_percent=40,
            message=f"Detected {len(detected_fields)} fields from PDF (Azure DI)"
        )

        payload["detection_result"] = detection_result
        payload["processing_time_ms"] = elapsed_ms
        return {"status": "completed", **payload}

    except SoftTimeLimitExceeded:
        logger.warning(f"detect_fields_task soft time limit exceeded for fill_run={fill_run_id}")
        repo.update_fill_run(fill_run_id, status="failed", error_stage="field_detection",
                             error_message="Task timed out (soft limit)")
        tracker.mark_error(error_stage="field_detection", error_message="Field detection timed out",
                           error_type="timeout", is_retryable=False)
        _reverse_template_fill_shadow(fill_run_id, "field_detection_timeout", "field_detection")
        db.close()
        return {"status": "failed", "error": "timeout", **payload}

    except Exception as e:
        logger.error(f"Field detection failed: {e}", exc_info=True)
        _fail_fill_run(repo, tracker, fill_run_id, "field_detection", e)
        _reverse_template_fill_shadow(fill_run_id, "field_detection_failed", "field_detection")
        db.close()
        return {"status": "failed", "error": str(e), **payload}


# ---------------------------------------------------------------------------
# Task: auto_map_fields_task
# ---------------------------------------------------------------------------

@shared_task(bind=True)
def auto_map_fields_task(self, payload: Dict[str, Any]) -> Dict[str, Any]:
    """
    Auto-map PDF fields to Excel cells using schema-first, then LLM fallback.

    Workflow:
    1. Try schema-based mapping (deterministic, instant)
    2. Targeted LLM for unmapped YAML fields
    3. Merge results (schema takes priority)
    """
    if payload.get("status") == "failed":
        logger.warning(f"Skipping auto-mapping because field detection failed: {payload.get('error')}")
        return payload

    fill_run_id = payload["fill_run_id"]
    template_id = payload["template_id"]
    detection_result = payload["detection_result"]
    job_id = payload["job_id"]

    skip_schema = payload.get("skip_schema", True)

    db = _get_db_session()
    repo = TemplateRepository(db)
    tracker = JobProgressTracker(db, job_id)

    _fill_run_check = repo.get_fill_run_with_data(fill_run_id)
    _AUTO_MAP_DONE_STATUSES = {"awaiting_review", "filling", "completed", "failed"}
    if _fill_run_check and _fill_run_check.status in _AUTO_MAP_DONE_STATUSES:
        logger.warning(
            f"auto_map_fields_task skipped — fill_run={fill_run_id} already in status={_fill_run_check.status!r}"
        )
        db.close()
        _idempotent_fm = _fill_run_check.fill_run_data.field_mapping if _fill_run_check.fill_run_data else None
        return {**payload, "mapping_result": _idempotent_fm or {}}

    try:
        task_start_time = time.time()
        logger.info(f"Auto-mapping fields for fill run: {fill_run_id}")

        tracker.update_progress(
            status="mapping",
            current_stage="auto_mapping",
            progress_percent=50,
            message="Mapping PDF fields to Excel cells"
        )

        template = repo.get_template(template_id)
        if not template:
            raise ValueError(f"Template not found: {template_id}")

        excel_schema = template.schema_metadata
        if not excel_schema:
            raise ValueError(f"Template schema not found: {template_id}")

        all_pdf_fields = list(detection_result.get("fields", []))
        pdf_fields = list(all_pdf_fields)

        fill_run_for_ctx = _fill_run_check
        citation_context = (fill_run_for_ctx.fill_run_data.citation_context or {"citations": []}) if (fill_run_for_ctx and fill_run_for_ctx.fill_run_data) else {"citations": []}
        context_budget: Dict[str, Any] = {
            "context_budget_applied": False,
            "user_warning": None,
            "scalar_batches": {},
            "table_batches": {},
        }
        schema_mappings = []
        schema_id = None
        schema_obj = None
        schema_table_row_labels: Dict[str, Any] = {}
        schema_summary: Optional[Dict[str, Any]] = None

        # === STEP 1: Identify template schema ===
        template_path = None
        workbook = None
        try:
            from openpyxl import load_workbook

            storage = get_storage_backend()
            template_path = storage.download_to_temp(template.file_path)
            workbook = load_workbook(template_path, data_only=False)

            schema_id = mapping_coordinator.identify_template(workbook)

            if schema_id:
                logger.info(f"✓ Template identified as: {schema_id}")
                schema_obj_for_rows = mapping_coordinator.schema_loader.load_schema(schema_id)
                if schema_obj_for_rows:
                    schema_obj = schema_obj_for_rows
                    schema_table_row_labels = extract_schema_table_row_labels(schema_obj_for_rows, workbook)
                    schema_summary = compute_schema_counts(schema_obj_for_rows)
                    logger.info(
                        "YAML schema summary: fields=%s table_cells=%s total=%s",
                        schema_summary.get("yaml_field_count"),
                        schema_summary.get("yaml_table_cells"),
                        schema_summary.get("total_yaml_fields"),
                    )

                if not skip_schema:
                    schema_mappings = mapping_coordinator.create_schema_mappings(schema_id, pdf_fields)
                    logger.info(f"Schema mapping: {len(schema_mappings)} fields mapped (confidence=0.98 baseline)")

                    try:
                        from app.verticals.real_estate.template_filling.excel.schema_based import SchemaMapper
                        schema = mapping_coordinator.schema_loader.load_schema(schema_id)
                        if schema:
                            mapper = SchemaMapper(schema)
                            unmapped_field_ids = mapper.get_unmapped_schema_fields(schema_mappings)
                            if unmapped_field_ids:
                                logger.warning(
                                    f"Schema fields without PDF data ({len(unmapped_field_ids)}): "
                                    f"{', '.join(unmapped_field_ids[:10])}"
                                    f"{'...' if len(unmapped_field_ids) > 10 else ''}"
                                )
                    except Exception as e:
                        logger.debug(f"Could not check unmapped schema fields: {e}")

                    tracker.update_progress(
                        status="mapping",
                        current_stage="auto_mapping",
                        progress_percent=50,
                        message=f"Schema matched {len(schema_mappings)} fields by alias"
                    )
                else:
                    logger.info("Skipping alias mapping (skip_schema=True) — will use LLM only for schema fields")
            else:
                raise ValueError(
                    f"Template {template_id} does not match any known schema. "
                    "Only pre-defined master templates are supported."
                )

        except ValueError:
            raise
        except Exception as e:
            raise RuntimeError(f"Schema identification failed for template {template_id}: {e}") from e
        finally:
            if workbook is not None:
                try:
                    workbook.close()
                except Exception:
                    pass
            if template_path:
                Path(template_path).unlink(missing_ok=True)

        # === STEP 1.5: Targeted LLM for unmapped YAML fields ===
        llm_service_targeted = None
        om_structure_artifact: Optional[Dict[str, Any]] = None
        structure_plan: Dict[str, Any] = {
            "fields_to_extract": [],
            "tables_to_extract": [],
            "skipped_targets": [],
            "review_required_targets": [],
        }
        if schema_id:
            try:
                from app.verticals.real_estate.template_filling.excel.schema_based import SchemaMapper

                schema_obj = mapping_coordinator.schema_loader.load_schema(schema_id)
                if schema_obj:
                    mapper_for_unmapped = SchemaMapper(schema_obj)
                    unmapped_field_ids = mapper_for_unmapped.get_unmapped_schema_fields(schema_mappings)
                    unmapped_fields = [
                        schema_obj.get_field(fid)
                        for fid in unmapped_field_ids
                        if schema_obj.get_field(fid) is not None
                    ]

                    schema_tables = schema_obj.tables or []
                    source_pdf_fields_for_table_context = list(all_pdf_fields)

                    if unmapped_fields or schema_tables:
                        llm_service_targeted = TemplateFillLLMService(
                            prompt_version=settings.re_template_prompt_version,
                            capture_io_log=settings.capture_llm_io_log,
                            fill_run_id=fill_run_id,
                        )

                        tracker.update_progress(
                            status="mapping",
                            current_stage="auto_mapping",
                            progress_percent=52,
                            message="Detecting OM source map..."
                        )
                        detected_structure = asyncio.run(
                            llm_service_targeted.detect_om_structure(all_pdf_fields)
                        )
                        detected_structure = normalize_om_structure_with_pdf_fields(
                            detected_structure,
                            all_pdf_fields,
                        )
                        om_structure_artifact = build_om_structure_artifact(
                            detected_structure,
                            llm_service_targeted.model,
                        )

                        repo.merge_fill_run_extracted_data(
                            fill_run_id,
                            {"om_structure": om_structure_artifact},
                        )

                        structure_plan = plan_schema_targets_for_structure(
                            unmapped_fields,
                            schema_tables,
                            detected_structure,
                        )
                        unmapped_fields = structure_plan["fields_to_extract"]
                        schema_tables = structure_plan["tables_to_extract"]
                        logger.info(
                            "OM structure plan: extract fields=%s tables=%s skipped=%s review=%s",
                            len(unmapped_fields),
                            len(schema_tables),
                            len(structure_plan["skipped_targets"]),
                            len(structure_plan["review_required_targets"]),
                        )

                    if unmapped_fields:
                        logger.info(f"🎯 Stage 2: {len(unmapped_fields)} unmapped YAML fields → targeted LLM")
                        tracker.update_progress(
                            status="mapping",
                            current_stage="auto_mapping",
                            progress_percent=55,
                            message=f"Finding values for {len(unmapped_fields)} schema fields..."
                        )

                        targeted_values: Dict[str, Any] = {}
                        scalar_batches: Dict[str, List[Dict[str, Any]]] = {}
                        for field_def in unmapped_fields:
                            scalar_batches.setdefault(get_scalar_batch_key(field_def), []).append(field_def)

                        prepared_scalar_batches: List[Dict[str, Any]] = []
                        for batch_key, batch_fields in scalar_batches.items():
                            scalar_context, scalar_budget = build_scalar_context_for_batch(
                                all_pdf_fields,
                                om_structure_artifact.get("effective") if om_structure_artifact else {},
                                batch_fields,
                            )
                            context_budget["scalar_batches"][batch_key] = scalar_budget
                            if scalar_budget.get("context_budget_applied"):
                                context_budget["context_budget_applied"] = True
                                if scalar_budget.get("user_warning"):
                                    context_budget["user_warning"] = scalar_budget["user_warning"]
                            logger.info(
                                "Scalar context prepared for batch %s: %s → %s fields (routing=%s, citation_pages=%s)",
                                batch_key,
                                scalar_budget["source_field_count_before_routing"],
                                len(scalar_context),
                                scalar_budget["routing_applied"],
                                scalar_budget["citation_pages"],
                            )
                            prepared_scalar_batches.append({
                                "batch_key": batch_key,
                                "fields": batch_fields,
                                "context": scalar_context,
                                "budget": scalar_budget,
                            })

                        consolidated_scalar_batches = consolidate_scalar_batches_by_context(
                            prepared_scalar_batches
                        )
                        context_budget["scalar_batch_summary"] = {
                            "original_batch_count": len(prepared_scalar_batches),
                            "consolidated_call_count": len(consolidated_scalar_batches),
                        }
                        logger.info(
                            "Scalar LLM call consolidation: original_batches=%s consolidated_calls=%s",
                            len(prepared_scalar_batches),
                            len(consolidated_scalar_batches),
                        )

                        for consolidated_batch in consolidated_scalar_batches:
                            batch_keys = consolidated_batch["batch_keys"]
                            batch_fields = consolidated_batch["fields"]
                            scalar_context = consolidated_batch["context"]
                            if len(batch_keys) > 1:
                                logger.info(
                                    "Merged scalar batches with identical context: batches=%s fields=%s",
                                    batch_keys,
                                    len(batch_fields),
                                )
                            batch_values = asyncio.run(
                                llm_service_targeted.extract_schema_field_values(
                                    batch_fields,
                                    scalar_context,
                                    om_structure=(
                                        om_structure_artifact.get("effective")
                                        if om_structure_artifact
                                        else None
                                    ),
                                )
                            )
                            targeted_values.update(batch_values)

                        targeted_mappings = mapping_coordinator.create_targeted_schema_mappings(
                            unmapped_fields, targeted_values
                        )

                        targeted_mapping_by_field_id = {
                            mapping.get("pdf_field_name"): mapping
                            for mapping in targeted_mappings
                        }
                        for field_def in unmapped_fields:
                            fid = field_def["id"]
                            mapping = targeted_mapping_by_field_id.get(fid)
                            if mapping:
                                pdf_fields.append(
                                    _build_targeted_virtual_pdf_field(mapping, citation_context)
                                )

                        schema_mappings = schema_mappings + targeted_mappings

                        self.update_state(
                            state="PROGRESS",
                            meta={"stage": "targeted_schema_mapping_done"}
                        )

                        tracker.update_progress(
                            status="mapping",
                            current_stage="auto_mapping",
                            progress_percent=70,
                            message=(
                                f"Found {len(targeted_mappings)} additional fields via targeted LLM "
                                f"({len(schema_mappings)} schema total)"
                            )
                        )

                    if schema_tables:
                        logger.info(f"🎯 Stage 2: {len(schema_tables)} schema tables → targeted LLM")
                        tracker.update_progress(
                            status="mapping",
                            current_stage="auto_mapping",
                            progress_percent=70,
                            message=f"Extracting values for {len(schema_tables)} schema tables..."
                        )

                        targeted_table_values = {}
                        table_batches: Dict[str, List[Dict[str, Any]]] = {}
                        for table_def in schema_tables:
                            table_id = table_def.get("id")
                            if not table_id:
                                continue
                            batch_key = get_table_batch_key(table_def)
                            table_batches.setdefault(batch_key, []).append(table_def)

                        effective_om = (
                            om_structure_artifact.get("effective") if om_structure_artifact else None
                        )

                        for batch_key, batch_tables in table_batches.items():
                            if not source_pdf_fields_for_table_context:
                                logger.info(f"No pdf_fields available for table batch {batch_key}")
                                continue

                            fill_when_keys: set = set()
                            for tbl in batch_tables:
                                for k in (tbl.get("fill_when") or []):
                                    fill_when_keys.add(k)
                            citation_pages: List[int] = []
                            for key in fill_when_keys:
                                citation_pages.extend(
                                    get_section_citation_pages(effective_om or {}, key)
                                )

                            if citation_pages:
                                filtered = [
                                    f for f in source_pdf_fields_for_table_context
                                    if (pg := get_field_page(f)) is None
                                    or any(abs(pg - p) <= 1 for p in citation_pages)
                                ]
                                logger.info(
                                    "Table context filtered for batch %s: %s → %s fields (citation pages %s)",
                                    batch_key, len(source_pdf_fields_for_table_context),
                                    len(filtered), sorted(set(citation_pages)),
                                )
                            else:
                                filtered = source_pdf_fields_for_table_context

                            table_context_fields, table_budget = build_budgeted_pdf_fields(filtered)
                            table_budget = {
                                **table_budget,
                                "routing_applied": bool(citation_pages),
                                "citation_pages": sorted(set(citation_pages)),
                                "source_field_count_before_routing": len(source_pdf_fields_for_table_context),
                                "source_field_count_after_routing": len(filtered),
                            }
                            context_budget["table_batches"][batch_key] = table_budget
                            if table_budget.get("context_budget_applied"):
                                context_budget["context_budget_applied"] = True
                                if table_budget.get("user_warning"):
                                    context_budget["user_warning"] = table_budget["user_warning"]
                            context_payload = llm_service_targeted._build_table_context_from_pdf_fields(table_context_fields)
                            table_context_json: Optional[str] = json.dumps(
                                context_payload, separators=(",", ":"), ensure_ascii=False
                            )

                            row_labels_by_table: Dict[str, List[str]] = {}
                            for table_def in batch_tables:
                                table_id = table_def.get("id")
                                if not table_id:
                                    continue
                                row_labels = schema_table_row_labels.get(table_id, {}).get("row_labels", [])
                                if row_labels and all(not str(lbl).strip() for lbl in row_labels):
                                    logger.info(f"Row labels empty for table {table_id}; using row_index order")
                                    row_labels = []
                                row_labels_by_table[table_id] = row_labels

                            logger.info(f"Table batch extraction: batch={batch_key} tables={len(batch_tables)}")

                            batch_result = asyncio.run(
                                llm_service_targeted.extract_schema_table_values_rag_batch(
                                    batch_tables,
                                    table_context_fields,
                                    row_labels_by_table=row_labels_by_table,
                                    prebuilt_context_json=table_context_json,
                                    om_structure=(
                                        om_structure_artifact.get("effective")
                                        if om_structure_artifact
                                        else None
                                    ),
                                )
                            )
                            if batch_result:
                                targeted_table_values.update(batch_result)

                        targeted_table_mappings = mapping_coordinator.create_targeted_schema_table_mappings(
                            schema_tables, targeted_table_values, schema_table_row_labels
                        )

                        for mapping in targeted_table_mappings:
                            extracted_value = mapping.get("extracted_value")
                            if extracted_value is None:
                                continue
                            pdf_fields.append(
                                _build_targeted_virtual_pdf_field(mapping, citation_context)
                            )

                        schema_mappings = schema_mappings + targeted_table_mappings

                        tracker.update_progress(
                            status="mapping",
                            current_stage="auto_mapping",
                            progress_percent=70,
                            message=(
                                f"Found {len(targeted_table_mappings)} table cells via targeted LLM "
                                f"({len(schema_mappings)} schema total)"
                            )
                        )

                    if not unmapped_fields and not schema_tables:
                        logger.info("All YAML schema fields matched in Stage 1 — skipping Stage 2")
                        tracker.update_progress(
                            status="mapping",
                            current_stage="auto_mapping",
                            progress_percent=70,
                            message=f"All schema fields matched ({len(schema_mappings)} total)"
                        )
            except Exception:
                logger.error("Stage 2 targeted mapping failed", exc_info=True)
                raise

        # === STEP 2: Merge schema mappings (schema alias + targeted LLM) ===
        raw_mappings = schema_mappings

        def _tier_priority(mapping: dict) -> int:
            source = mapping.get("source")
            if source == "schema":
                return 2
            if source == "targeted_schema":
                return 1
            return 0

        best_by_cell: dict[str, dict] = {}
        for m in raw_mappings:
            excel_sheet = m.get("excel_sheet", "")
            excel_cell = m.get("excel_cell", "")
            if not excel_sheet or not excel_cell:
                continue

            cell_key = f"{excel_sheet}!{excel_cell}"
            current_best = best_by_cell.get(cell_key)
            current_best_tier = _tier_priority(current_best or {})
            candidate_tier = _tier_priority(m)
            current_best_conf = float((current_best or {}).get("confidence") or 0)
            candidate_conf = float(m.get("confidence") or 0)

            if (
                current_best is None
                or candidate_tier > current_best_tier
                or (candidate_tier == current_best_tier and candidate_conf > current_best_conf)
            ):
                best_by_cell[cell_key] = m

        mappings = list(best_by_cell.values())
        total_mapped_fields = len(mappings)

        unique_pdf_fields_mapped = len(set(m.get("pdf_field_id") for m in mappings if m.get("pdf_field_id")))
        targeted_count = sum(1 for m in mappings if m.get("source") == "targeted_schema")
        schema_count = sum(1 for m in mappings if m.get("source") == "schema")

        logger.info(
            f"Mapping deduplication: {len(raw_mappings)} raw → "
            f"{len(mappings)} after cell dedup "
            f"({unique_pdf_fields_mapped} unique PDF fields mapped to {len(mappings)} Excel cells) "
            f"[{schema_count} alias + {targeted_count} targeted_llm]"
        )

        mapping_result = {
            "mappings": mappings,
            "total_mapped": total_mapped_fields,
            "schema_mapped_count": schema_count,
            "targeted_schema_count": targeted_count,
            "high_confidence_count": sum(1 for m in mappings if m.get("confidence", 0) >= STRUCTURE_HIGH_CONFIDENCE)
        }

        field_mapping = {
            "pdf_fields": pdf_fields,
            "mappings": mappings,
        }
        if context_budget.get("context_budget_applied"):
            field_mapping["context_budget"] = context_budget
        if schema_obj:
            field_mapping["yaml_cell_status"] = build_yaml_cell_status(
                schema_obj,
                mappings,
                skipped_targets=structure_plan.get("skipped_targets"),
                review_required_targets=structure_plan.get("review_required_targets"),
            )
        if schema_summary:
            field_mapping["schema_summary"] = schema_summary
        if om_structure_artifact:
            field_mapping["om_structure_summary"] = {
                "extractor_version": om_structure_artifact.get("extractor_version"),
                "confidence_summary": om_structure_artifact.get("confidence_summary"),
                "edited": om_structure_artifact.get("edited"),
            }

        metadata_params = {
            "field_mapping": field_mapping,
            "total_fields_mapped": total_mapped_fields,
            "auto_mapped_count": total_mapped_fields,
            "user_edited_count": 0,
            "auto_mapping_completed": True,
        }

        metadata_params["status"] = "awaiting_review"
        metadata_params["current_stage"] = "auto_mapping"
        map_elapsed_ms = int((time.time() - task_start_time) * 1000)
        metadata_params["processing_time_ms"] = payload.get("processing_time_ms", 0) + map_elapsed_ms
        blob_params = {k: metadata_params.pop(k) for k in ["field_mapping", "citation_context", "extracted_data"] if k in metadata_params}
        if blob_params:
            repo.update_fill_run_data(fill_run_id, **blob_params)
        repo.update_fill_run(fill_run_id, **metadata_params)

        db.close()

        logger.info(
            f"Auto-mapping complete: {len(mappings)} Excel cells mapped "
            f"(from {unique_pdf_fields_mapped} unique PDF fields) — "
            f"{schema_count} alias + {targeted_count} targeted LLM"
        )

        parts = []
        if schema_count:
            parts.append(f"{schema_count} alias")
        if targeted_count:
            parts.append(f"{targeted_count} targeted LLM")
        status_msg = (
            f"Mapped {len(mappings)} cells ({', '.join(parts)})"
            if parts
            else f"Mapped {len(mappings)} cells"
        )

        tracker.update_progress(
            status="awaiting_review",
            current_stage="auto_mapping",
            progress_percent=85,
            message=status_msg
        )

        payload["mapping_result"] = mapping_result
        payload["processing_time_ms"] = payload.get("processing_time_ms", 0) + map_elapsed_ms
        return {"status": "completed", **payload}

    except SoftTimeLimitExceeded:
        logger.warning(f"auto_map_fields_task soft time limit exceeded for fill_run={fill_run_id}")
        repo.update_fill_run(fill_run_id, status="failed", error_stage="auto_mapping",
                             error_message="Task timed out (soft limit)")
        tracker.mark_error(error_stage="auto_mapping", error_message="Auto-mapping timed out",
                           error_type="timeout", is_retryable=False)
        _reverse_template_fill_shadow(fill_run_id, "auto_mapping_timeout", "auto_mapping")
        db.close()
        return {"status": "failed", "error": "timeout", **payload}

    except Exception as e:
        logger.error(f"Auto-mapping failed: {e}", exc_info=True)
        _mark_auto_mapping_exception(repo, tracker, fill_run_id, e)
        _reverse_template_fill_shadow(fill_run_id, "auto_mapping_failed", "auto_mapping")
        db.close()
        return {"status": "failed", "error": str(e), **payload}


# ---------------------------------------------------------------------------
# Task: fill_excel_task
# ---------------------------------------------------------------------------

@shared_task(bind=True)
def fill_excel_task(self, payload: Dict[str, Any]) -> Dict[str, Any]:
    """Fill Excel template with data from Azure DI (extracted_value)."""
    if payload.get("status") == "failed":
        logger.warning(f"Skipping Excel filling because previous task failed: {payload.get('error')}")
        return payload

    fill_run_id = payload["fill_run_id"]
    template_id = payload["template_id"]
    job_id = payload["job_id"]

    db = _get_db_session()
    repo = TemplateRepository(db)
    tracker = JobProgressTracker(db, job_id)

    _fill_run_check = repo.get_fill_run(fill_run_id)
    if _fill_run_check and _fill_run_check.status in ("completed", "failed"):
        logger.warning(
            f"fill_excel_task skipped — fill_run={fill_run_id} already in status={_fill_run_check.status!r}"
        )
        db.close()
        return {**payload, "status": _fill_run_check.status}

    start_time = time.monotonic()

    try:
        logger.info(f"Filling Excel for fill run: {fill_run_id}")

        tracker.update_progress(
            status="filling",
            current_stage="excel_filling",
            progress_percent=90,
            message="Filling Excel template"
        )

        fill_run = repo.get_fill_run_with_data(fill_run_id)
        if not fill_run:
            raise ValueError(f"Fill run not found: {fill_run_id}")

        field_mapping = fill_run.fill_run_data.field_mapping if fill_run.fill_run_data else {}

        template = repo.get_template(template_id)
        if not template:
            raise ValueError(f"Template not found: {template_id}")

        file_ext = template.file_extension or ".xlsx"

        storage = get_storage_backend()
        template_local_path = f"/tmp/template_{template_id}{file_ext}"
        storage.download(template.file_path, template_local_path)

        output_local_path = f"/tmp/filled_{fill_run_id}{file_ext}"

        extracted_data = _prepare_extracted_data_for_fill(
            fill_run.fill_run_data.extracted_data if fill_run.fill_run_data else {},
            field_mapping,
        )
        llm_extracted = extracted_data.get("llm_extracted", {})
        manual_edits = extracted_data.get("manual_edits", {})

        user_edited_count = sum(len(cells) for cells in manual_edits.values())
        auto_mapped_count = len(llm_extracted) - sum(1 for data in llm_extracted.values() if data.get('user_edited'))

        logger.info(
            f"Prepared extracted_data for filling: {user_edited_count} user-edited + "
            f"{auto_mapped_count} LLM auto-mapped = {len(llm_extracted)} total"
        )

        handler = ExcelHandler()
        fill_summary = handler.fill_template(
            template_path=template_local_path,
            output_path=output_local_path,
            field_mapping=field_mapping,
            extracted_data=extracted_data
        )

        _now = datetime.now(timezone.utc)
        safe_template_name = template.name.replace("/", "_").replace("\\", "_")
        storage_key = (
            f"fills/{_now.year}/{_now.month:02d}/{_now.day:02d}"
            f"/{fill_run_id}_{safe_template_name}_filled{file_ext}"
        )
        storage.upload(output_local_path, storage_key)

        artifact = {
            "backend": settings.storage_backend,
            "key": storage_key,
            "size": Path(output_local_path).stat().st_size,
            "filename": f"{safe_template_name}_filled{file_ext}"
        }

        repo.update_fill_run(
            fill_run_id,
            artifact=artifact,
            status="completed",
            filling_completed=True,
            completed_at=datetime.now(timezone.utc),
            total_fields_filled=fill_summary.get("total_cells_filled", 0),
            processing_time_ms=(fill_run.processing_time_ms or 0) + int((time.monotonic() - start_time) * 1000)
        )

        record_template_fill_completed(org_id=fill_run.org_id)

        latency = time.monotonic() - start_time
        try:
            TEMPLATE_FILL_LATENCY_SECONDS.observe(latency)
        except Exception as e:
            logger.warning(f"Failed to record template fill latency: {e}", exc_info=True)

        db.close()

        logger.info(f"Excel filling complete: {fill_summary['total_cells_filled']} cells filled")
        commit_shadow_credits(operation_type="template_fill_run", reference_id=fill_run_id)

        tracker.mark_completed()

        payload["artifact"] = artifact
        payload["fill_summary"] = fill_summary
        return {"status": "completed", **payload}

    except SoftTimeLimitExceeded:
        logger.warning(f"fill_excel_task soft time limit exceeded for fill_run={fill_run_id}")
        repo.update_fill_run(fill_run_id, status="failed", error_stage="excel_filling",
                             error_message="Task timed out (soft limit)")
        tracker.mark_error(error_stage="excel_filling", error_message="Excel filling timed out",
                           error_type="timeout", is_retryable=False)
        _reverse_template_fill_shadow(fill_run_id, "excel_filling_timeout", "excel_filling")
        db.close()
        return {"status": "failed", "error": "timeout", **payload}

    except Exception as e:
        logger.error(f"Excel filling failed: {e}", exc_info=True)

        repo.update_fill_run(
            fill_run_id,
            status="failed",
            error_stage="excel_filling",
            error_message=str(e)
        )

        record_template_fill_failed(org_id=fill_run.org_id if fill_run else None)

        latency = time.monotonic() - start_time
        try:
            TEMPLATE_FILL_LATENCY_SECONDS.observe(latency)
        except Exception as e:
            logger.warning(f"Failed to record template fill latency: {e}", exc_info=True)

        tracker.mark_error(
            error_stage="excel_filling",
            error_message="Failed to fill template — please try again.",
            internal_error=str(e)[:1000],
            error_type="fill_error",
            is_retryable=False
        )
        _reverse_template_fill_shadow(fill_run_id, "excel_filling_failed", "excel_filling")

        db.close()
        return {"status": "failed", "error": str(e), **payload}


# ---------------------------------------------------------------------------
# Task: start_fill_run_chain
# ---------------------------------------------------------------------------

@shared_task(bind=True)
def start_fill_run_chain(
    self,
    template_id: str,
    document_id: str,
    user_id: str,
    fill_run_name: str | None = None,
) -> str:
    """Orchestrate the full template filling pipeline."""
    db = _get_db_session()
    repo = TemplateRepository(db)

    try:
        logger.info(f"Starting fill run chain: template={template_id}, document={document_id}")

        template = repo.get_template(template_id)
        if not template:
            raise ValueError(f"Template not found: {template_id}")

        template_snapshot = {
            "name": template.name,
            "description": template.description,
            "schema_metadata": template.schema_metadata,
        }

        fill_run = repo.create_fill_run(
            template_id=template_id,
            document_id=document_id,
            org_id=template.org_id,
            user_id=user_id,
            template_snapshot=template_snapshot,
            name=fill_run_name,
        )

        fill_run_id = fill_run.id
        job_id = fill_run_id
        job_repo = JobRepository()
        job_state = job_repo.create_job(
            entity_type="template_fill_run",
            entity_id=fill_run_id,
            status="queued",
            current_stage="initialization",
            progress_percent=5,
            job_id=job_id
        )
        if not job_state:
            raise ValueError("Failed to create job state for fill run")

        db.close()

        task_chain = chain(
            detect_fields_task.s({
                "fill_run_id": fill_run_id,
                "template_id": template_id,
                "document_id": document_id,
                "job_id": job_id,
            }),
            auto_map_fields_task.s(),
        )

        task_chain.apply_async(queue='critical')

        logger.info(f"Fill run chain started: {fill_run_id} (queue=critical)")

        return fill_run_id

    except Exception as e:
        logger.error(f"Failed to start fill run chain: {e}", exc_info=True)
        db.close()
        raise


# ---------------------------------------------------------------------------
# Task: continue_fill_run_chain
# ---------------------------------------------------------------------------

@shared_task(bind=True)
def continue_fill_run_chain(
    self,
    fill_run_id: str,
    job_id: str,
) -> Dict[str, Any]:
    """Continue fill run pipeline after user has reviewed mappings."""
    db = _get_db_session()
    repo = TemplateRepository(db)
    tracker = JobProgressTracker(db, job_id)

    try:
        logger.info(f"Continuing fill run: {fill_run_id}")

        fill_run = repo.get_fill_run(fill_run_id)
        if not fill_run:
            raise ValueError(f"Fill run not found: {fill_run_id}")

        repo.update_fill_run(
            fill_run_id,
            user_review_completed=True,
            status="filling"
        )

        db.close()

        fill_excel_task.apply_async(
            kwargs={
                "payload": {
                    "fill_run_id": fill_run_id,
                    "template_id": fill_run.template_id,
                    "job_id": job_id,
                }
            }
        )

        logger.info(f"Fill run continuation started (direct to filling): {fill_run_id}")

        return {"fill_run_id": fill_run_id, "status": "processing"}

    except Exception as e:
        logger.error(f"Failed to continue fill run: {e}", exc_info=True)

        tracker.update_progress(
            status="failed",
            current_stage="continuation",
            message=f"Failed to continue fill run: {str(e)}"
        )
        _reverse_template_fill_shadow(fill_run_id, "continue_fill_failed", "continuation")

        db.close()
        raise
